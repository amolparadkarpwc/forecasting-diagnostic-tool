"""Planted-scenario tests for the MVP diagnostic."""

from __future__ import annotations

import json
import os
import pickle
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
PACK = ROOT / "pack"
DEMO = ROOT / "examples" / "demo_input"

sys.path.insert(0, str(ROOT))

from diag.blocks.fact_pattern import run as fact_pattern  # noqa: E402
from diag.blocks.phrase_search import run as phrase_search  # noqa: E402
from diag.handlers.documents import parse_document  # noqa: E402
from diag.intake.identify import identify  # noqa: E402
from diag.intake.safety import assess_model  # noqa: E402
from diag.intake.unpack import extract_zip  # noqa: E402
from diag.rules.derived import rate_above  # noqa: E402
from diag.runner import init_engagement, report_engagement, run_engagement  # noqa: E402


def _scores(results: dict) -> dict:
    return {item["id"]: item for item in results["diagnostics"]}


def _png_1x1(red: int = 255) -> bytes:
    import struct
    import zlib

    def chunk(tag: bytes, payload: bytes) -> bytes:
        crc = zlib.crc32(tag + payload) & 0xFFFFFFFF
        return struct.pack(">I", len(payload)) + tag + payload + struct.pack(">I", crc)

    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    raw = bytes((0, red, 0, 0))
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")


def _record(path: str, detected: str) -> dict:
    return {
        "artifact_id": "A001",
        "path": path,
        "detected_type": detected,
        "family": "document",
        "encoding": "",
        "status": "readable",
        "reason": "",
        "period_covered": "",
        "version_group": Path(path).stem.lower(),
    }


def _dashboard_payload(html: str) -> dict:
    marker = '<script id="results" type="application/json">'
    start = html.index(marker) + len(marker)
    end = html.index("</script>", start)
    return json.loads(html[start:end])


class DiagnosticTests(unittest.TestCase):
    def test_help(self) -> None:
        completed = subprocess.run(
            [sys.executable, "-m", "diag", "--help"],
            cwd=ROOT,
            capture_output=True,
            text=True,
            check=False,
        )
        self.assertEqual(completed.returncode, 0, completed.stderr)
        self.assertIn("finalise", completed.stdout)
        self.assertIn("run", completed.stdout)

    def test_pack_and_derived_helper(self) -> None:
        from diag.packload import load_pack

        pack = load_pack(PACK)
        self.assertEqual(pack["meta"]["version"], "0.5.0")
        self.assertEqual({item["id"] for item in pack["diagnostics"]}, {"D08", "D43"})
        self.assertTrue(rate_above(0.3, 0.2))
        self.assertFalse(rate_above(0.1, 0.2))

    def test_source_does_not_unpickle_or_execute(self) -> None:
        banned = ("pickle.loads", "pickle.load(", "exec(", "eval(")
        text = "\n".join(path.read_text(encoding="utf-8") for path in (ROOT / "diag").rglob("*.py"))
        for token in banned:
            self.assertNotIn(token, text)

    def test_zip_slip_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            archive = folder / "evil.zip"
            dest = folder / "out"
            with zipfile.ZipFile(archive, "w") as handle:
                handle.writestr("ok.txt", "hello")
                handle.writestr("../escaped.txt", "nope")
                handle.writestr("..\\also.txt", "nope")
            extracted, rejected = extract_zip(archive, dest)
            self.assertEqual([item["member"] for item in extracted], ["ok.txt"])
            self.assertTrue(any("escapes" in item["reason"] for item in rejected))
            self.assertFalse((folder / "escaped.txt").exists())
            self.assertFalse(any(path.name in {"escaped.txt", "also.txt"} for path in dest.rglob("*")))

    def test_disallowed_pickle_is_quarantined_and_not_loaded(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            marker = folder / "PWNED.txt"
            engagement = folder / "eng"

            class Evil:
                def __reduce__(self):
                    return (os.system, (f'echo pwned > "{marker}"',))

            init_engagement(engagement)
            (engagement / "input" / "bad.pkl").write_bytes(pickle.dumps(Evil(), protocol=4))
            run_engagement(engagement, PACK)
            self.assertFalse(marker.exists())
            inventory = json.loads((engagement / "output" / "inventory.json").read_text(encoding="utf-8"))
            pickle_row = next(row for row in inventory if row["path"].endswith("bad.pkl"))
            self.assertEqual(pickle_row["status"], "quarantined")
            self.assertIn("disallowed", pickle_row["reason"])
            status, reason = assess_model(engagement / "input" / "bad.pkl", "pickle", ["sklearn"])
            self.assertEqual(status, "quarantined")
            self.assertTrue("os" in reason or "nt" in reason or "posix" in reason)

    def test_label_shift_is_not_scored_as_leakage(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            engagement = Path(tmp) / "eng"
            init_engagement(engagement)
            shutil.copy(ROOT / "tests" / "planted" / "label_shift" / "train.py", engagement / "input" / "train.py")
            results = run_engagement(engagement, PACK)
            evidence = json.loads((engagement / "output" / "evidence_log.json").read_text(encoding="utf-8"))
            leakage = [
                record
                for record in evidence["records"]
                if record["signal_id"] == "negative_shift" and record["status"] in {"auto_accepted", "counted_flagged"}
            ]
            review = [record for record in evidence["records"] if record["signal_id"] == "possible_label_shift"]
            self.assertEqual(leakage, [])
            self.assertIn("D08", _scores(results))
            self.assertNotIn("D14", _scores(results))

    def test_naive_fallback_mention_needs_review(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            engagement = Path(tmp) / "eng"
            init_engagement(engagement)
            shutil.copy(ROOT / "tests" / "planted" / "naive_fallback" / "note.md", engagement / "input" / "note.md")
            results = run_engagement(engagement, PACK)
            self.assertIn("D08", _scores(results))
            self.assertNotIn("D34", _scores(results))

    def test_demo_end_to_end(self) -> None:
        self.assertTrue(DEMO.exists(), "run python examples/build_demo_input.py first")
        with tempfile.TemporaryDirectory() as tmp:
            engagement = Path(tmp) / "demo"
            init_engagement(engagement)
            shutil.copytree(DEMO, engagement / "input", dirs_exist_ok=True)
            results = run_engagement(engagement, PACK)
            scores = _scores(results)
            self.assertEqual(set(scores), {"D08", "D43"})
            self.assertEqual(results["versions"]["pack"], "0.5.0")
            self.assertIn(
                results["overall"]["verdict"],
                {"Ready", "Ready with gaps", "Not ready", "Readiness not confirmed"},
            )
            self.assertIn("verdict_reason", results["overall"])
            d43 = scores["D43"]
            self.assertIn(d43.get("gate_result", {}).get("status"), {"passes", "tripped", "pending", "no_evidence", None})
            self.assertIn("python", results["versions"])
            inventory = json.loads((engagement / "output" / "inventory.json").read_text(encoding="utf-8"))
            self.assertTrue(any(row["status"] == "quarantined" for row in inventory))
            self.assertTrue(any(row["status"] == "duplicate" for row in inventory))
            self.assertTrue(any(row["status"] == "unreadable" for row in inventory))
            onnx = next(row for row in inventory if row["detected_type"] == "onnx")
            self.assertEqual(onnx["status"], "manual_review")
            self.assertIn("ONNX", onnx["reason"])
            pdf = next(row for row in inventory if row["detected_type"] == "pdf")
            self.assertEqual(pdf["status"], "manual_review")
            output_text = "\n".join(
                path.read_text(encoding="utf-8") for path in (engagement / "output").iterdir() if path.suffix != ".html"
            )
            self.assertNotIn("AcmeForecast", output_text)
            self.assertNotIn("Jane Doe", output_text)
            self.assertIn("[REDACTED]", output_text)
            dashboard = report_engagement(engagement)
            html = dashboard.read_text(encoding="utf-8")
            self.assertIn(results["overall"]["verdict"], html)
            self.assertIn("Summary", html)
            self.assertIn("Coverage", html)
            self.assertIn("Scorecard", html)
            self.assertIn("Recommendations", html)
            self.assertIn("What-if", html)
            self.assertIn("Measured evidence only", html)
            self.assertNotIn("AcmeForecast", html)
            self.assertNotIn("Client A", html)
            self.assertNotIn("carefully guarded", html)
            self.assertNotIn("fonts.googleapis.com", html)
            self.assertNotIn("client_pipeline.py", html)
            payload = _dashboard_payload(html)
            by_id = {item["id"]: item for item in payload["checks"]}
            self.assertEqual(payload["engine"]["final"]["verdict"], results["overall"]["verdict"])
            self.assertEqual(payload["engine"]["final"]["score"], results["overall"]["score"])
            self.assertEqual(by_id["D43"]["score"], scores["D43"]["score"])
            self.assertTrue(by_id["D43"]["inPack"])
            self.assertIsNone(by_id["D21"]["score"])
            self.assertEqual(by_id["D21"]["ev"], "NE")
            self.assertFalse(by_id["D21"]["inPack"])
            self.assertEqual(by_id["D21"]["name"], "Experiment tracking")
            self.assertIsNone(by_id["D22"]["score"])
            self.assertIsNone(by_id["D10"]["score"])
            self.assertNotEqual(by_id["D14"]["score"], 82)

    def test_repeat_run_matches(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            first = folder / "one"
            second = folder / "two"
            for engagement in (first, second):
                init_engagement(engagement)
                config_path = engagement / "engagement.yaml"
                config_text = config_path.read_text(encoding="utf-8")
                config_path.write_text(
                    config_text.replace(f"engagement_id: {engagement.name}", "engagement_id: demo"),
                    encoding="utf-8",
                )
                shutil.copytree(DEMO, engagement / "input", dirs_exist_ok=True)
                run_engagement(engagement, PACK)
            left = json.loads((first / "output" / "results.json").read_text(encoding="utf-8"))
            right = json.loads((second / "output" / "results.json").read_text(encoding="utf-8"))
            left["versions"].pop("generated_at", None)
            right["versions"].pop("generated_at", None)
            self.assertEqual(left, right)

    def test_vtt_transcript_evidence_points_at_cue(self) -> None:
        fixture = ROOT / "tests" / "planted" / "transcripts" / "meeting.vtt"
        artifact = parse_document(fixture, _record("meeting.vtt", "vtt"))
        self.assertIsNotNone(artifact)
        assert artifact is not None
        self.assertNotIn("-->", artifact.text)
        self.assertNotIn("align:", artifact.text)
        self.assertNotIn("WEBVTT", artifact.text)
        self.assertNotIn("generated for the diagnostic fixture", artifact.text)
        self.assertEqual(artifact.sections[0]["speaker"], "Priya Shah")
        self.assertEqual(artifact.sections[0]["timestamp"], "00:00:10.000")
        self.assertEqual(artifact.sections[0]["location"], {"kind": "cue", "value": "00:00:10.000"})
        self.assertIn("naive baseline", artifact.sections[0]["text"])
        self.assertNotIn("naive baseline", artifact.sections[1]["text"])
        self.assertNotIn("Priya Shah", artifact.text)
        phrases = phrase_search(artifact, {"phrases": ["naive baseline"]}, {}, {})
        self.assertEqual(phrases[0].location, {"kind": "cue", "value": "00:00:10.000"})
        facts = fact_pattern(artifact, {"pattern": "naive baseline"}, {}, {})
        self.assertEqual(facts[0].location, {"kind": "cue", "value": "00:00:10.000"})
        with tempfile.TemporaryDirectory() as tmp:
            engagement = Path(tmp) / "eng"
            init_engagement(engagement)
            shutil.copy(fixture, engagement / "input" / "meeting.vtt")
            run_engagement(engagement, PACK)
            evidence = json.loads((engagement / "output" / "evidence_log.json").read_text(encoding="utf-8"))
            hits = [
                record
                for record in evidence["records"]
                if str(record["source"]["file"]).endswith("meeting.vtt")
            ]
            self.assertTrue(all(record["source"]["location"]["kind"] == "cue" for record in hits))
            self.assertTrue(all(record["evidence_level"] == "Attested" for record in hits))
            self.assertTrue(all(record["handler"] == "documents" for record in hits))

    def test_webvtt_text_and_speaker_lines_are_transcripts(self) -> None:
        fixture = ROOT / "tests" / "planted" / "transcripts" / "meeting.txt"
        detected, family, _encoding = identify(fixture, fixture.read_bytes())
        self.assertEqual(detected, "vtt")
        self.assertEqual(family, "document")
        self.assertNotIn(detected, {"unknown", "txt"})
        speaker = b"Jane Doe 0:01:02\nWe compared against a naive baseline.\n"
        detected, family, _encoding = identify(Path("talk.txt"), speaker)
        self.assertEqual((detected, family), ("transcript", "document"))
        voiced = b"<v Jane Doe>hello from the meeting</v>\n"
        detected, family, _encoding = identify(Path("voice.txt"), voiced)
        self.assertEqual((detected, family), ("transcript", "document"))
        detected, family, _encoding = identify(Path("notes.vtt"), b"not a caption header\n")
        self.assertEqual((detected, family), ("vtt", "document"))
        with tempfile.TemporaryDirectory() as tmp:
            engagement = Path(tmp) / "eng"
            init_engagement(engagement)
            shutil.copy(fixture, engagement / "input" / "notes.txt")
            talk = engagement / "input" / "talk.txt"
            talk.write_bytes(speaker)
            run_engagement(engagement, PACK)
            inventory = json.loads((engagement / "output" / "inventory.json").read_text(encoding="utf-8"))
            notes = next(row for row in inventory if row["path"].endswith("notes.txt"))
            self.assertEqual(notes["detected_type"], "vtt")
            self.assertEqual(notes["family"], "document")
            self.assertEqual(notes["status"], "readable")
            self.assertNotEqual(notes["detected_type"], "unknown")
            spoken_row = next(row for row in inventory if row["path"].endswith("talk.txt"))
            self.assertEqual(spoken_row["detected_type"], "transcript")
            self.assertEqual(spoken_row["family"], "document")
            self.assertEqual(spoken_row["status"], "readable")
            spoken = parse_document(talk, _record("talk.txt", "transcript"))
            assert spoken is not None
            self.assertEqual(spoken.sections[0]["speaker"], "Jane Doe")
            self.assertEqual(spoken.sections[0]["location"], {"kind": "cue", "value": "0:01:02"})
            self.assertNotIn("0:01:02", spoken.text)
            matches = phrase_search(spoken, {"phrases": ["naive baseline"]}, {}, {})
            self.assertEqual(matches[0].location, {"kind": "cue", "value": "0:01:02"})

    def test_png_is_inventoried_without_requiring_tesseract(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            engagement = Path(tmp) / "eng"
            init_engagement(engagement)
            (engagement / "input" / "chart.png").write_bytes(_png_1x1(255))
            (engagement / "input" / "renamed.txt").write_bytes(_png_1x1(1))
            run_engagement(engagement, PACK)
            inventory = json.loads((engagement / "output" / "inventory.json").read_text(encoding="utf-8"))
            installed = shutil.which("tesseract") is not None
            for name in ("chart.png", "renamed.txt"):
                row = next(item for item in inventory if item["path"].endswith(name))
                self.assertEqual(row["detected_type"], "png")
                self.assertEqual(row["family"], "document")
                self.assertNotEqual(row["status"], "unreadable")
                if installed:
                    self.assertIn(row["status"], {"readable", "manual_review"})
                else:
                    self.assertEqual(row["status"], "manual_review")
                    self.assertIn("OCR not installed", row["reason"])

    def test_corrupt_image_does_not_abort_the_run(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            engagement = Path(tmp) / "eng"
            init_engagement(engagement)
            (engagement / "input" / "broken.png").write_bytes(b"\x89PNG\r\n\x1a\n" + b"this is not a png")
            (engagement / "input" / "note.txt").write_text(
                "A naive baseline is the comparison.\n",
                encoding="utf-8",
            )
            results = run_engagement(engagement, PACK)
            self.assertIn("diagnostics", results)
            inventory = json.loads((engagement / "output" / "inventory.json").read_text(encoding="utf-8"))
            broken = next(row for row in inventory if row["path"].endswith("broken.png"))
            self.assertEqual(broken["detected_type"], "png")
            self.assertEqual(broken["family"], "document")
            self.assertEqual(broken["status"], "unreadable")
            self.assertIn("corrupt", broken["reason"])
            note = next(row for row in inventory if row["path"].endswith("note.txt"))
            self.assertEqual(note["status"], "readable")

    def test_image_signatures(self) -> None:
        webp_body = b"WEBP" + b"VP8 " + (4).to_bytes(4, "little") + b"\x00\x00\x00\x00"
        samples = {
            "jpeg": b"\xff\xd8\xff\xe0" + b"\x00" * 8 + b"\xff\xd9",
            "gif": b"GIF89a" + b"\x01\x00\x01\x00\x00\x00\x00" + b"\x3b",
            "webp": b"RIFF" + len(webp_body).to_bytes(4, "little") + webp_body,
            "tiff": b"II*\x00" + (8).to_bytes(4, "little") + b"\x00\x00",
        }
        names = {"jpeg": "photo.jpg", "gif": "frame.gif", "webp": "shot.webp", "tiff": "scan.tiff"}
        for kind, data in samples.items():
            detected, family, _encoding = identify(Path(names[kind]), data)
            self.assertEqual(detected, kind)
            self.assertEqual(family, "document")
        gif87, family, _encoding = identify(Path("old.gif"), b"GIF87a" + b"\x01\x00\x01\x00\x80\x00\x00" + b"\x3b")
        self.assertEqual((gif87, family), ("gif", "document"))
        big, family, _encoding = identify(Path("scan.tif"), b"MM\x00*" + (8).to_bytes(4, "big") + b"\x00\x00")
        self.assertEqual((big, family), ("tiff", "document"))
        renamed, family, _encoding = identify(Path("notes.txt"), _png_1x1())
        self.assertEqual((renamed, family), ("png", "document"))
        labeled, family, _encoding = identify(Path("notes.png"), b"hello from a mislabeled note\n")
        self.assertEqual(labeled, "txt")
        self.assertNotEqual(labeled, "png")

    def test_image_ocr_text_is_attested(self) -> None:
        from types import SimpleNamespace
        from unittest.mock import patch

        fake = SimpleNamespace(returncode=0, stdout=b"The slide cites a naive baseline.\n", stderr=b"")
        with tempfile.TemporaryDirectory() as tmp:
            engagement = Path(tmp) / "eng"
            init_engagement(engagement)
            (engagement / "input" / "figure.png").write_bytes(_png_1x1())
            with (
                patch("diag.handlers.images.shutil.which", return_value="tesseract"),
                patch("diag.handlers.images.subprocess.run", return_value=fake),
            ):
                run_engagement(engagement, PACK)
            inventory = json.loads((engagement / "output" / "inventory.json").read_text(encoding="utf-8"))
            row = next(item for item in inventory if item["path"].endswith("figure.png"))
            self.assertEqual(row["detected_type"], "png")
            self.assertEqual(row["status"], "readable")
            evidence = json.loads((engagement / "output" / "evidence_log.json").read_text(encoding="utf-8"))
            hits = [
                record
                for record in evidence["records"]
                if str(record["source"]["file"]).endswith("figure.png")
            ]
            self.assertTrue(all(record["evidence_level"] == "Attested" for record in hits))
            self.assertTrue(all(record["handler"] == "documents" for record in hits))

    def test_docx_transcript_cues(self) -> None:
        try:
            import docx
        except ImportError:
            self.skipTest("python-docx is not installed")
        with tempfile.TemporaryDirectory() as tmp:
            engagement = Path(tmp) / "eng"
            init_engagement(engagement)
            document = docx.Document()
            document.add_paragraph("Priya Shah 0:04:05")
            document.add_paragraph("The notes mention a naive baseline.")
            document.save(engagement / "input" / "transcript.docx")
            run_engagement(engagement, PACK)
            inventory = json.loads((engagement / "output" / "inventory.json").read_text(encoding="utf-8"))
            row = next(item for item in inventory if item["path"].endswith("transcript.docx"))
            self.assertEqual(row["detected_type"], "docx")
            self.assertEqual(row["family"], "document")
            self.assertEqual(row["status"], "readable")
            evidence = json.loads((engagement / "output" / "evidence_log.json").read_text(encoding="utf-8"))
            hits = [
                record
                for record in evidence["records"]
                if str(record["source"]["file"]).endswith("transcript.docx")
            ]
            self.assertTrue(all(record["evidence_level"] == "Attested" for record in hits))


if __name__ == "__main__":
    unittest.main()
