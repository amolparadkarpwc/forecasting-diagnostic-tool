"""Rebuild examples/demo_input. Safe to re-run. Does not execute client code or load pickles."""

from __future__ import annotations

import csv
import json
import os
import pickle
import zipfile
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "examples" / "demo_input"

SEASON = [100, 110, 105, 120, 130, 180, 170, 160, 150, 140, 125, 200]
SKIP = {("A", 2023, 4), ("A", 2024, 9), ("B", 2024, 11)}


def actual(series: str, year: int, month: int) -> int:
    base = SEASON[month - 1]
    if series == "A":
        return base + (25 if year == 2024 else 0)
    return base * 2 + (40 if year == 2024 else 0)


def write(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text.replace("\r\n", "\n"), encoding="utf-8", newline="\n")


def build_tables() -> None:
    sales_path = OUT / "data" / "sales_history.csv"
    forecast_path = OUT / "data" / "past_forecasts.csv"
    sales_path.parent.mkdir(parents=True, exist_ok=True)
    sales_rows = []
    forecast_rows = []
    for series in ("A", "B"):
        for year in (2023, 2024):
            for month in range(1, 13):
                if (series, year, month) in SKIP:
                    continue
                value = actual(series, year, month)
                sales_rows.append({"series_id": series, "date": date(year, month, 1).isoformat(), "actual": value})
                if year == 2024:
                    if month == 1:
                        origin = date(2023, 12, 1)
                    else:
                        origin = date(2024, month - 1, 1)
                    error = 8 if series == "A" else 12
                    forecast_rows.append(
                        {
                            "series_id": series,
                            "run_date": origin.isoformat(),
                            "target_date": date(year, month, 1).isoformat(),
                            "forecast": value - error,
                        }
                    )
    sales_rows.append(
        {"series_id": "A", "date": "2024-01-01", "actual": actual("A", 2024, 1)}
    )
    with sales_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["series_id", "date", "actual"])
        writer.writeheader()
        writer.writerows(sales_rows)
    with forecast_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["series_id", "run_date", "target_date", "forecast"])
        writer.writeheader()
        writer.writerows(forecast_rows)


def build_code() -> None:
    write(
        OUT / "code" / "features.py",
        '''"""Feature build for the demand model."""
import warnings

import pandas as pd

warnings.filterwarnings("ignore")


def build_features(df):
    """future_sales looks ahead of the row and is a leakage risk."""
    grouped = df.groupby("series_id")["actual"]
    df["future_sales"] = grouped.shift(-1)
    df["lag_1"] = grouped.shift(1)
    return df


def fit_models(models, features, target):
    for model in models:
        try:
            model.fit(features, target)
        except Exception:
            continue
''',
    )
    write(
        OUT / "code" / "features.sql",
        """SELECT s.series_id, s.week_start, p.price
FROM sales s
JOIN price_actuals p ON p.week_start >= s.week_start
""",
    )
    write(OUT / "code" / "broken.py", "def (\n")
    notebook = {
        "nbformat": 4,
        "nbformat_minor": 5,
        "metadata": {},
        "cells": [
            {"cell_type": "markdown", "metadata": {}, "source": ["Training notes\n"]},
            {
                "cell_type": "code",
                "metadata": {},
                "execution_count": None,
                "outputs": [],
                "source": ["import pandas as pd\n", "train = df.copy()\n"],
            },
        ],
    }
    write(OUT / "code" / "train.ipynb", json.dumps(notebook, indent=2) + "\n")


def build_docs() -> None:
    write(
        OUT / "docs" / "methodology.md",
        """# Forecast methodology

We measure accuracy with WAPE and bias for AcmeForecast (owner Jane Doe) at horizons of 1 and 3 months.
The production model is compared with a seasonal naive baseline.
Validation checks run before each forecast.
""",
    )
    write(
        OUT / "config" / "settings.yaml",
        """client: AcmeForecast
owner: Jane Doe
metrics:
  primary: WAPE
  secondary: bias
backtest:
  horizons: [1, 3]
model:
  family: gradient_boosting
  n_estimators: 200
""",
    )
    write(
        OUT / "logs" / "pipeline.log",
        """2024-01-15T02:00:00Z job=forecast status=SUCCESS duration=120
2024-02-15T02:00:01Z job=forecast status=FAILED duration=12 error=timeout
2024-03-15T02:00:02Z job=forecast status=SUCCESS duration=118
2024-04-15T02:00:03Z job=forecast status=FAILED duration=8 error=data_missing
""",
    )
    write(OUT / "notes.txt", "Runbook excerpt for the monthly forecast.\n")
    write(OUT / "notes_copy.txt", "Runbook excerpt for the monthly forecast.\n")
    write(OUT / "broken.json", "{ this is not json\n")
    (OUT / "docs").mkdir(parents=True, exist_ok=True)
    (OUT / "docs" / "brief.pdf").write_bytes(b"%PDF-1.4\n1 0 obj\n<<>>\nendobj\ntrailer\n<<>>\n%%EOF\n")
    (OUT / "models").mkdir(parents=True, exist_ok=True)
    (OUT / "models" / "candidate.onnx").write_bytes(b"onnx-mvp-placeholder\n")

    from docx import Document

    document = Document()
    document.add_heading("Model card", level=1)
    document.add_paragraph("Primary metric is WAPE for AcmeForecast, prepared by Jane Doe.")
    document.save(OUT / "docs" / "model_card.docx")

    class Marker:
        def __reduce__(self):
            return (os.system, ("echo diagnostic-mvp",))

    (OUT / "models" / "suspicious.pkl").write_bytes(pickle.dumps(Marker(), protocol=4))

    with zipfile.ZipFile(OUT / "bundle.zip", "w") as handle:
        handle.writestr("supplement.txt", "Scheduler note: the monthly forecast job is forecast_monthly.\n")


def main() -> None:
    if OUT.exists():
        import shutil

        shutil.rmtree(OUT)
    build_tables()
    build_code()
    build_docs()
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
